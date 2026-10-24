"""Synthetic mathematical fixtures only; no aircraft performance data."""
import copy
import importlib.util
import json
import math
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from engine.technical_reader_v1 import Axis, ReaderError, TechnicalStore, evaluate, interpolate, linearize, parse_number, validate_model
from engine.technical_reader_v1.ingest import guided_table, ingest
from engine.technical_reader_v1.models import digitize_trace
from engine.technical_reader_v1.numeric import convert_unit
from engine.technical_reader_v1.plans import calculate_plan
from engine.technical_reader_v1.tables import table_to_grid, describe_table
from engine.technical_reader_v1.raster import scanline_candidates, check_model


def model(kind="curve"):
    base = {"id": "SYNTHETIC-MATH", "revision": 1, "kind": kind, "title": "Synthetic test function",
            "quality": "validated", "validation": {"evidence": ["analytical mathematical fixture"]},
            "provenance": [{"source_id": "TEST", "source_revision": "r1", "pdf_page": 1}],
            "inputs": [{"name": "x", "unit": "1"}], "output": {"name": "y", "unit": "1"}}
    if kind == "curve":
        base["segments"] = [{"points": [[0, 0], [5, 10], [10, 0]]}]
    elif kind == "family":
        base["inputs"].append({"name": "h", "unit": "1"})
        base["curves"] = [{"id": "low", "parameter": 0, "segments": [{"points": [[0, 0], [10, 20]]}]},
                          {"id": "high", "parameter": 10, "segments": [{"points": [[0, 10], [10, 30]]}]}]
    elif kind == "grid":
        base["inputs"].append({"name": "z", "unit": "1"})
        base["axes"] = [{"name": "x", "values": [0, 10]}, {"name": "z", "values": [0, 10]}]
        base["values"] = [0, 30, 20, 50]  # y=2*x+3*z, last axis varies fastest
    elif kind == "lookup":
        base["inputs"] = [{"name": "mode", "unit": "category", "type": "category"}]
        base["output"]["unit"] = "text"
        base["rows"] = [{"keys": ["A"], "value": "normal"}, {"keys": ["B"], "value": "alternate"}]
    else:
        base["expression"] = "2*x + 1"
    return base


class NumericTests(unittest.TestCase):
    def test_regular_numbers_and_persian(self):
        for raw, value in [("۱۲٫۵", 12.5), ("−2", -2), ("1e-2", .01), (".5", .5)]:
            self.assertEqual(parse_number(raw)["value"], value)

    def test_locale_is_explicit(self):
        self.assertEqual(parse_number("1,234")["kind"], "ambiguous")
        self.assertEqual(parse_number("1,234", thousands=",")["value"], 1234)
        self.assertEqual(parse_number("12,5", decimal=",")["value"], 12.5)
        self.assertEqual(parse_number("1,2", thousands=",")["kind"], "ambiguous")

    def test_blank_not_zero(self):
        for raw in (None, "", "—", "N/A"):
            self.assertEqual(parse_number(raw)["kind"], "missing")
            self.assertIsNone(parse_number(raw)["value"])

    def test_bounds_ranges_and_ocr_ambiguity(self):
        self.assertEqual(parse_number("≤ 15")["kind"], "bound")
        self.assertEqual(parse_number("-2 to 4")["minimum"], -2)
        self.assertEqual(parse_number("O.5")["kind"], "text")

    def test_nonfinite_rejected(self):
        with self.assertRaises(ReaderError):
            interpolate([0, 1], [0, float("nan")], .5)

    def test_linear_and_knots(self):
        self.assertEqual(interpolate([0, 1, 2], [0, 4, 1], 1)[0], 4)
        self.assertEqual(interpolate([0, 2], [0, 4], 1)[0], 2)

    def test_extrapolation_rejected(self):
        with self.assertRaises(ReaderError):
            interpolate([0, 1], [0, 2], 1.1)

    def test_duplicate_x_rejected(self):
        with self.assertRaises(ReaderError):
            interpolate([0, 0], [0, 1], 0)

    def test_axis_reversed_and_roundtrip(self):
        axis = Axis(((10, 100), (110, 0)), unit="psi")
        self.assertEqual(axis.to_value(60), 50)
        self.assertEqual(axis.to_pixel(50), 60)
        self.assertEqual(axis.interval(60, 2), [48, 52])

    def test_axis_piecewise(self):
        self.assertEqual(Axis(((0, 0), (10, 1), (20, 3))).to_value(15), 2)

    def test_axis_log(self):
        axis = Axis(((0, 1), (100, 100)), scale="log10")
        self.assertAlmostEqual(axis.to_value(50), 10)
        self.assertAlmostEqual(axis.to_pixel(10), 50)

    def test_axis_invalid(self):
        for anchors in [((0, 0), (0, 1)), ((0, 0), (1, 1), (2, 0))]:
            with self.assertRaises(ReaderError):
                Axis(anchors)

    def test_unit_conversions(self):
        self.assertAlmostEqual(convert_unit(1000, "ft", "m"), 304.8)
        self.assertAlmostEqual(convert_unit(32, "degF", "degC"), 0)
        with self.assertRaises(ReaderError):
            convert_unit(10, "kg", "psi")


class ModelTests(unittest.TestCase):
    def test_curve_kink(self):
        self.assertEqual(evaluate(model(), {"x": 5})["value"], 10)
        self.assertEqual(evaluate(model(), {"x": 7.5})["value"], 5)

    def test_family(self):
        result = evaluate(model("family"), {"x": 3, "h": 5})
        self.assertEqual(result["value"], 11)
        self.assertEqual(len(result["trace"]["curves"]), 2)

    def test_family_exact_level_does_not_touch_invalid_adjacent_domain(self):
        m = model("family")
        m["curves"][1]["segments"][0]["points"] = [[5, 20], [10, 30]]
        self.assertEqual(evaluate(m, {"x": 1, "h": 0})["value"], 2)
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 1, "h": 5})

    def test_grid(self):
        self.assertEqual(evaluate(model("grid"), {"x": 2, "z": 3})["value"], 13)

    def test_nd_grid(self):
        m = model("grid")
        m["inputs"].append({"name": "t", "unit": "1"})
        m["axes"].append({"name": "t", "values": [0, 1]})
        m["values"] = [0, 1, 30, 31, 20, 21, 50, 51]
        self.assertAlmostEqual(evaluate(m, {"x": 2, "z": 3, "t": .5})["value"], 13.5)

    def test_grid_missing_corner_rejected_but_exact_other_cell_works(self):
        m = model("grid"); m["values"][-1] = None
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 1, "z": 1})
        self.assertEqual(evaluate(m, {"x": 0, "z": 0})["value"], 0)

    def test_grid_shape(self):
        m = model("grid"); m["values"].pop()
        with self.assertRaises(ReaderError):
            validate_model(m)

    def test_lookup_text(self):
        self.assertEqual(evaluate(model("lookup"), {"mode": "B"})["value"], "alternate")
        with self.assertRaises(ReaderError):
            evaluate(model("lookup"), {"mode": "C"})

    def test_formula(self):
        self.assertEqual(evaluate(model("formula"), {"x": 3})["value"], 7)

    def test_formula_code_rejected(self):
        for expr in ["__import__('os').getcwd()", "x.__class__", "[x for x in range(2)]", "x**100"]:
            m = model("formula"); m["expression"] = expr
            with self.assertRaises(ReaderError):
                evaluate(m, {"x": 1})

    def test_formula_division_by_zero(self):
        m = model("formula"); m["expression"] = "1/x"
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 0})

    def test_conditions_do_not_fallback(self):
        m = model(); m["conditions"] = {"engine": "VARIANT-A", "bleed": "on"}
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 1}, context={"engine": "VARIANT-B", "bleed": "on"})
        self.assertEqual(evaluate(m, {"x": 1}, context=m["conditions"])["value"], 2)

    def test_quality_gate_and_explicit_approximation(self):
        for quality in ("candidate", "rejected"):
            m = model(); m["quality"] = quality
            with self.assertRaises(ReaderError):
                evaluate(m, {"x": 1}, mode="simulator")
        m = model(); m["quality"] = "simulation_approx"
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 1})
        self.assertEqual(evaluate(m, {"x": 1}, mode="simulator")["quality"], "simulation_approx")

    def test_source_required(self):
        m = model(); m["provenance"][0]["pdf_page"] = 0
        with self.assertRaises(ReaderError):
            validate_model(m)

    def test_error_bound_not_invented(self):
        m = model()
        self.assertIsNone(evaluate(m, {"x": 1})["interval"])
        m["absolute_error"] = .2
        self.assertEqual(evaluate(m, {"x": 1})["interval"], [1.8, 2.2])

    def test_gaps_and_ambiguous_branches(self):
        m = model(); m["segments"] = [{"points": [[0, 0], [1, 1]]}, {"points": [[2, 2], [3, 3]]}]
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 1.5})
        m["segments"][1]["points"][0][0] = .5
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": .75})

    def test_linearize_preserves_knots_and_gaps(self):
        m = model(); m["segments"] = [{"points": [[0, 0], [.3, 2], [1, 1]]}, {"points": [[2, 2], [3, 3]]}]
        rows = linearize(m, .5)
        self.assertTrue(any(r["x"] == .3 for r in rows))
        self.assertFalse(any(1 < r["x"] < 2 for r in rows))

    def test_log_y_curve(self):
        m = model(); m["segments"] = [{"points": [[0, 1], [10, 100]]}]; m["y_scale"] = "log10"
        self.assertAlmostEqual(evaluate(m, {"x": 5})["value"], 10)

    def test_digitize_pixels(self):
        seg = digitize_trace([[[10, 90], [60, 40], [110, 90]]], Axis(((10, 0), (110, 10))), Axis(((40, 10), (90, 0))))
        self.assertEqual(seg[0]["points"], [[0, 0], [5, 10], [10, 0]])


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = TechnicalStore(Path(self.tmp.name) / "test.sqlite")
        self.store.source("TEST", "r1", "0" * 64, {"synthetic": True})
        self.store.add_page("TEST", "r1", {"pdf_page": 1, "width": 100, "height": 100, "coordinate_space": "pixel", "method": "synthetic", "spans": [{"text": "sample", "bbox": [0, 0, 10, 10], "line_id": [1]}, {"text": "pressure", "bbox": [10, 0, 20, 10], "line_id": [1]}], "tables": []})

    def tearDown(self):
        self.store.close(); self.tmp.cleanup()

    def test_model_storage_and_calculation_audit(self):
        self.store.add_model(model())
        r = self.store.calculate("SYNTHETIC-MATH", 1, {"x": 1})
        self.assertEqual(r["value"], 2)
        self.assertEqual(self.store.conn.execute("SELECT count(*) FROM ctr_curve_points").fetchone()[0], 3)
        self.assertEqual(self.store.conn.execute("SELECT count(*) FROM ctr_calculations").fetchone()[0], 1)

    def test_immutable_revision_and_idempotency(self):
        m = model(); digest = self.store.add_model(m)
        self.assertEqual(self.store.add_model(m), digest)
        m["segments"][0]["points"][0][1] = 2
        with self.assertRaises(ReaderError):
            self.store.add_model(m)
        m["revision"] = 2; self.store.add_model(m)
        self.assertEqual(self.store.get_model(m["id"], 1)["segments"][0]["points"][0][1], 0)

    def test_missing_source_rolls_back_model(self):
        m = model(); m["provenance"][0]["source_id"] = "MISSING"
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.add_model(m)
        self.assertEqual(self.store.conn.execute("SELECT count(*) FROM ctr_models").fetchone()[0], 0)

    def test_source_revision_conflict(self):
        with self.assertRaises(ReaderError):
            self.store.source("TEST", "r1", "1" * 64, {"synthetic": True})

    def test_phrase_search(self):
        self.assertEqual(self.store.search("sample pressure")[0]["pdf_page"], 1)

    def test_chart_region_inventory(self):
        self.store.region("REGION-1", "TEST", "r1", 1, "chart", "NEEDS_TRACE", {"bbox": [0, 0, 10, 10]})
        row = self.store.conn.execute("SELECT * FROM ctr_regions").fetchone()
        self.assertEqual(row["status"], "NEEDS_TRACE")
        self.assertEqual(json.loads(row["metadata_json"])["bbox"], [0, 0, 10, 10])

    def test_export_integrity(self):
        self.assertIn("ctr_sources", self.store.export()["tables"])
        self.assertEqual(self.store.conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertEqual(self.store.conn.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_csv_end_to_end(self):
        path = Path(self.tmp.name) / "data.csv"
        path.write_text('name,value,blank\nA,12,\nB,,\nC,"1,234",\n')
        r = ingest(self.store, path, "CSV", "r1")
        self.assertEqual(r["table_candidates"], 1)
        kinds = [row[0] for row in self.store.conn.execute("SELECT value_kind FROM ctr_numeric_cells WHERE col_index=1 ORDER BY row_index")]
        self.assertEqual(kinds, ["text", "number", "missing", "ambiguous"])

    def test_plan_corrections(self):
        a = model("formula"); a["id"] = "BASE"; a["expression"] = "2*x"
        b = model("formula"); b["id"] = "CORRECTED"; b["expression"] = "base + delta"
        b["inputs"] = [{"name": "base", "unit": "1"}, {"name": "delta", "unit": "1"}]
        self.store.add_model(a); self.store.add_model(b)
        steps = [{"id": "a", "model_id": "BASE", "revision": 1, "bindings": {"x": {"input": "x"}}},
                 {"id": "b", "model_id": "CORRECTED", "revision": 1, "bindings": {"base": {"step": "a"}, "delta": {"constant": -1}}, "expected_units": {"a": "1"}}]
        self.assertEqual(calculate_plan(self.store, steps, {"x": 3})["value"], 5)
        steps[1]["expected_units"]["a"] = "psi"
        with self.assertRaises(ReaderError):
            calculate_plan(self.store, steps, {"x": 3})

    def test_cli(self):
        proc = subprocess.run([sys.executable, "-m", "engine.technical_reader_v1", "--db", str(self.store.path), "search", "sample pressure"], cwd=Path(__file__).resolve().parents[2], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertEqual(json.loads(proc.stdout)[0]["source_id"], "TEST")


class ExtractionTests(unittest.TestCase):
    def test_table_to_grid_requires_explicit_mapping_and_keeps_quality(self):
        m = table_to_grid([["x/z", "0", "10"], ["0", "0", "30"], ["10", "20", "50"]], model("grid"))
        self.assertEqual(m["quality"], "candidate")
        self.assertEqual(m["values"], [0, 30, 20, 50])
        with self.assertRaises(ReaderError):
            evaluate(m, {"x": 2, "z": 3}, mode="simulator")

    def test_table_bound_not_scalar_and_missing_preserved(self):
        rows = [["x/z", "0", "10"], ["0", "", "30"], ["10", "20", "50"]]
        self.assertIsNone(table_to_grid(rows, model("grid"))["values"][0])
        rows[1][1] = "<10"
        with self.assertRaises(ReaderError):
            table_to_grid(rows, model("grid"))

    def test_table_text_output(self):
        text = describe_table({"rows": [["x", "y"], [1, None]]}, {"source_id": "TEST", "source_revision": "r1", "pdf_page": 2})
        self.assertIn("C1=[MISSING]", text)
        self.assertIn("page: 2", text)

    def test_holdout_qa_does_not_promote_quality(self):
        m = model(); m["quality"] = "traced"
        report = check_model(m, [{"inputs": {"x": 2}, "expected": 4}], tolerance=.01)
        self.assertTrue(report["passed"])
        self.assertFalse(report["quality_promoted"])
        self.assertEqual(m["quality"], "traced")

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow optional")
    def test_raster_crossing_candidates_not_assigned_labels(self):
        from PIL import Image, ImageDraw
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "raster.png"
            image = Image.new("RGB", (100, 100), "white")
            draw = ImageDraw.Draw(image); draw.line((0, 20, 99, 20), fill="black"); draw.line((0, 80, 99, 80), fill="black")
            image.save(path)
            result = scanline_candidates(path, [0, 0, 100, 100], step=20)
            self.assertEqual(result["branch_assignment"], "unresolved")
            self.assertTrue(all(c["ambiguous"] for c in result["columns"]))

    def test_guided_table_retains_blank_cells(self):
        table = guided_table([{"text": "12", "bbox": [1, 1, 9, 9]}, {"text": "34", "bbox": [11, 11, 19, 19]}], [0, 10, 20], [0, 10, 20])
        self.assertEqual(table["rows"], [["12", ""], ["", "34"]])

    @unittest.skipUnless(importlib.util.find_spec("fitz"), "PyMuPDF optional")
    def test_pdf_native_physical_page_and_vectors(self):
        import fitz
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "source.pdf"
            doc = fitz.open(); doc.new_page()
            page = doc.new_page(); page.insert_text((30, 30), "SYNTHETIC PRESSURE 42")
            page.draw_line((30, 50), (100, 100)); doc.save(path); doc.close()
            store = TechnicalStore(Path(temp) / "test.sqlite")
            try:
                result = ingest(store, path, "PDF", "r1", pages=[2])
                self.assertEqual(result["pages"], 1)
                row = store.conn.execute("SELECT * FROM ctr_pages").fetchone()
                self.assertEqual(row["pdf_page"], 2)
                payload = json.loads(row["payload_json"])
                self.assertTrue(payload["vector_paths"])
                self.assertEqual(payload["coordinate_space"], "pdf_point")
                self.assertTrue(store.search("PRESSURE 42"))
            finally:
                store.close()

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl optional")
    def test_xlsx_preserves_formulas_and_merged_cells(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "source.xlsx"
            book = openpyxl.Workbook(); sheet = book.active
            sheet["A1"] = "Title"; sheet.merge_cells("A1:B1"); sheet["A2"] = 12; sheet["B2"] = "=A2*2"
            book.save(path); book.close()
            store = TechnicalStore(Path(temp) / "test.sqlite")
            try:
                ingest(store, path, "XLSX", "r1")
                payload = json.loads(store.conn.execute("SELECT payload_json FROM ctr_pages").fetchone()[0])
                self.assertEqual(payload["tables"][0]["merged_ranges"], ["A1:B1"])
                self.assertEqual(payload["tables"][0]["rows"][1][1], "=A2*2")
            finally:
                store.close()

    @unittest.skipUnless(shutil.which("tesseract") and importlib.util.find_spec("PIL"), "Tesseract and Pillow optional")
    def test_image_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "source.png"
            image = Image.new("RGB", (700, 130), "white")
            ImageDraw.Draw(image).text((20, 35), "PRESSURE 42", fill="black", font=ImageFont.truetype("DejaVuSans.ttf", 48))
            image.save(path)
            store = TechnicalStore(Path(temp) / "test.sqlite")
            try:
                result = ingest(store, path, "IMAGE", "r1", psm=6)
                self.assertGreater(result["spans"], 0)
                self.assertTrue(store.search("PRESSURE 42"))
                self.assertEqual(result["numeric_chart_models_automatically_validated"], 0)
            finally:
                store.close()

    @unittest.skipUnless(shutil.which("tesseract") and importlib.util.find_spec("PIL") and importlib.util.find_spec("fitz"), "PDF OCR dependencies optional")
    def test_scanned_pdf_ocr_has_physical_page_and_pdf_point_boxes(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        with tempfile.TemporaryDirectory() as temp:
            image_path = Path(temp) / "scan.png"
            pdf_path = Path(temp) / "scan.pdf"
            image = Image.new("RGB", (800, 180), "white")
            ImageDraw.Draw(image).text((20, 30), "TEMPERATURE 25", fill="black", font=ImageFont.truetype("DejaVuSans.ttf", 48))
            image.save(image_path)
            doc = fitz.open(); page = doc.new_page(width=400, height=90)
            page.insert_image(page.rect, filename=str(image_path)); doc.save(pdf_path); doc.close()
            store = TechnicalStore(Path(temp) / "test.sqlite")
            try:
                ingest(store, pdf_path, "SCANNED", "r1", dpi=144)
                row = store.conn.execute("SELECT * FROM ctr_pages").fetchone()
                self.assertEqual(row["extraction_method"], "tesseract_pdf_render")
                payload = json.loads(row["payload_json"])
                self.assertTrue(store.search("TEMPERATURE 25"))
                self.assertTrue(all(0 <= s["bbox"][0] < s["bbox"][2] <= 400 for s in payload["spans"]))
            finally:
                store.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
