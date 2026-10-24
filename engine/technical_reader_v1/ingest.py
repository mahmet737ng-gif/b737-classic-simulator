"""Native PDF, scanned PDF/image OCR, CSV/TSV/XLSX and guided table extraction.

Automatic output is always a candidate. OCR cannot identify a curve's altitude
label at an intersection; chart tracing is a separate explicit stage.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import subprocess
from pathlib import Path

from .numeric import ReaderError, finite


def source_hash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def _ocr(png_bytes, lang="eng", psm=3):
    try:
        proc = subprocess.run(["tesseract", "stdin", "stdout", "-l", lang, "--psm", str(psm), "tsv"],
                              input=png_bytes, capture_output=True, timeout=180, check=True)
    except FileNotFoundError as exc:
        raise ReaderError("OCR requires local Tesseract; no network OCR service is called") from exc
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise ReaderError("OCR failed; source has not been marked extracted") from exc
    spans = []
    for row in csv.DictReader(io.StringIO(proc.stdout.decode("utf-8")), delimiter="\t"):
        text = row.get("text", "").strip()
        if row.get("level") != "5" or not text:
            continue
        x, y, w, h = [int(row[n]) for n in ("left", "top", "width", "height")]
        spans.append({"text": text, "bbox": [x, y, x + w, y + h],
                      "confidence": float(row["conf"]) / 100,
                      "line_id": [int(row[n]) for n in ("block_num", "par_num", "line_num")]})
    return spans


def guided_table(spans, row_edges, col_edges):
    """Use declared row/column borders; do not collapse blank rows or merged cells."""
    for edges in (row_edges, col_edges):
        if len(edges) < 2 or any(finite(b) <= finite(a) for a, b in zip(edges, edges[1:])):
            raise ReaderError("Table boundaries must be increasing")
    rows = [[[] for _ in range(len(col_edges) - 1)] for _ in range(len(row_edges) - 1)]
    warnings = []
    for span in spans:
        x0, y0, x1, y1 = span["bbox"]
        x, y = (x0 + x1) / 2, (y0 + y1) / 2
        r = next((i for i, (a, b) in enumerate(zip(row_edges, row_edges[1:])) if a <= y < b), None)
        c = next((i for i, (a, b) in enumerate(zip(col_edges, col_edges[1:])) if a <= x < b), None)
        if r is None or c is None:
            continue
        rows[r][c].append(span)
        if x0 < col_edges[c] or x1 > col_edges[c + 1] or y0 < row_edges[r] or y1 > row_edges[r + 1]:
            warnings.append({"row": r, "col": c, "text": span["text"], "reason": "word_crosses_cell_border"})
    return {"rows": [[" ".join(s["text"] for s in sorted(cell, key=lambda s: (s["bbox"][1], s["bbox"][0]))) for cell in row] for row in rows],
            "row_edges": row_edges, "col_edges": col_edges, "warnings": warnings,
            "status": "candidate", "method": "guided_cell_boundaries", "merged_cells_inferred": False}


def extract_image(path, *, lang="eng", psm=3, crop=None):
    from PIL import Image, ImageOps
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        width, height = image.size
        if crop is not None:
            if len(crop) != 4 or not (0 <= crop[0] < crop[2] <= width and 0 <= crop[1] < crop[3] <= height):
                raise ReaderError("Crop lies outside the source image")
            image = image.crop(crop)
        buf = io.BytesIO()
        image.save(buf, format="PNG")
    spans = _ocr(buf.getvalue(), lang=lang, psm=psm)
    if crop is not None:
        for s in spans:
            x0, y0, x1, y1 = s["bbox"]
            s["bbox"] = [x0 + crop[0], y0 + crop[1], x1 + crop[0], y1 + crop[1]]
    return [{"pdf_page": 1, "width": width, "height": height, "coordinate_space": "pixel",
             "method": "tesseract_tsv", "status": "candidate", "crop": crop,
             "spans": spans, "tables": [], "chart_trace_status": "not_traced"}]


def extract_pdf(path, *, pages=None, dpi=300, lang="eng", force_ocr=False):
    import fitz
    if not 72 <= dpi <= 1200:
        raise ReaderError("DPI must be 72..1200")
    output = []
    with fitz.open(path) as doc:
        selected = list(range(1, len(doc) + 1)) if pages is None else list(pages)
        if not selected or len(selected) != len(set(selected)) or any(type(p) is not int or not 1 <= p <= len(doc) for p in selected):
            raise ReaderError("Invalid or duplicate physical PDF pages")
        for n in selected:
            page = doc[n - 1]
            words = page.get_text("words")
            warnings = []
            if words and not force_ocr:
                spans = [{"text": w[4], "bbox": list(w[:4]), "confidence": None, "line_id": list(w[5:7])} for w in words]
                method, space = "native_pdf_words", "pdf_point"
            else:
                pix = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
                spans = _ocr(pix.tobytes("png"), lang=lang)
                # Store OCR boxes in the same PDF-point coordinate system as tables.
                for s in spans:
                    s["bbox"] = [v * 72 / dpi for v in s["bbox"]]
                method, space = "tesseract_pdf_render", "pdf_point"
            tables = []
            if not force_ocr and words and hasattr(page, "find_tables"):
                try:
                    for table in page.find_tables().tables:
                        tables.append({"rows": table.extract(), "bbox": list(table.bbox),
                                       "method": "native_pdf_table_detector", "status": "candidate"})
                except Exception as exc:
                    warnings.append(f"table_detection_failed:{type(exc).__name__}")
            vectors = []
            for drawing in page.get_drawings():
                items = []
                for item in drawing["items"]:
                    if item[0] == "l":
                        items.append({"kind": "line", "points": [list(item[1]), list(item[2])]})
                    elif item[0] == "c":
                        items.append({"kind": "bezier", "points": [list(p) for p in item[1:]]})
                if items:
                    vectors.append({"items": items, "bbox": list(drawing["rect"]), "width": drawing.get("width"), "color": drawing.get("color")})
            output.append({"pdf_page": n, "width": page.rect.width, "height": page.rect.height,
                           "coordinate_space": space, "method": method, "status": "candidate", "spans": spans,
                           "tables": tables, "vector_paths": vectors, "warnings": warnings,
                           "chart_trace_status": "vector_candidates_only" if vectors else "not_traced"})
    return output


def extract_delimited(path, *, delimiter=None, numeric_locale=None):
    path = Path(path)
    delimiter = delimiter or ("\t" if path.suffix.lower() == ".tsv" else ",")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=delimiter))
    cols = max(map(len, rows), default=0)
    # Preserve ragged shape in metadata and distinguish absent cells from explicit blanks.
    metadata = {"original_row_lengths": list(map(len, rows)), "numeric_locale": numeric_locale or {}}
    padded = [row + [None] * (cols - len(row)) for row in rows]
    return [{"pdf_page": 1, "width": cols, "height": len(rows), "coordinate_space": "table_index",
             "method": "delimited_text", "status": "candidate", "spans": [],
             "tables": [{"rows": padded, "method": "csv_reader", "status": "candidate", **metadata}]}]


def extract_xlsx(path):
    import openpyxl
    formula_book = openpyxl.load_workbook(path, data_only=False)
    value_book = openpyxl.load_workbook(path, data_only=True)
    result = []
    try:
        for index, sheet in enumerate(formula_book.worksheets, 1):
            rows, cells = [], {}
            for row in sheet.iter_rows():
                values = []
                for cell in row:
                    raw = cell.value
                    values.append(None if raw is None else str(raw))
                    if raw is not None:
                        cells[f"{cell.row - 1}:{cell.column - 1}"] = {"address": cell.coordinate, "number_format": cell.number_format,
                            "formula": raw if cell.data_type == "f" else None,
                            "cached_value": str(value_book[sheet.title][cell.coordinate].value) if cell.data_type == "f" else None}
                rows.append(values)
            result.append({"pdf_page": index, "width": sheet.max_column, "height": sheet.max_row,
                           "coordinate_space": "table_index", "method": "xlsx_cells", "status": "candidate", "spans": [],
                           "tables": [{"rows": rows, "sheet": sheet.title, "cell_metadata": cells,
                                       "merged_ranges": [str(r) for r in sheet.merged_cells.ranges], "status": "candidate",
                                       "formula_execution": "not_performed"}]})
    finally:
        formula_book.close(); value_book.close()
    return result


def ingest(store, path, source_id, revision, *, metadata=None, **options):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = extract_pdf(path, **options)
    elif suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
        pages = extract_image(path, **options)
    elif suffix in (".csv", ".tsv"):
        pages = extract_delimited(path, **options)
    elif suffix == ".xlsx":
        pages = extract_xlsx(path)
    elif suffix in (".txt", ".md"):
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        pages = [{"pdf_page": 1, "width": max(map(len, lines), default=0), "height": len(lines),
                  "coordinate_space": "text_index", "method": "utf8_text", "status": "candidate", "tables": [],
                  "spans": [{"text": line, "bbox": [0, i, max(len(line), 1), i + 1], "confidence": None} for i, line in enumerate(lines)]}]
    else:
        raise ReaderError(f"Unsupported file format: {suffix}")
    digest = source_hash(path)
    store.source(source_id, revision, digest, {"filename": path.name, "source_type": suffix,
                 "privacy": "private_source", **(metadata or {})})
    for page in pages:
        store.add_page(source_id, revision, page)
    return {"source_id": source_id, "revision": revision, "sha256": digest, "pages": len(pages),
            "spans": sum(len(p["spans"]) for p in pages), "table_candidates": sum(len(p["tables"]) for p in pages),
            "quality": "candidate", "numeric_chart_models_automatically_validated": 0}
